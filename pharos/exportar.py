"""Exportação: Markdown, EPUB, DOCX e HTML bilíngue (EPUB escrito com zipfile, sem dependências)."""
from __future__ import annotations

import html
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from .config import Config, nome_idioma
from .modelo import Bloco, Livro, LEGENDA, NOTA, SEPARADOR, TITULO, VERSO, livro_para_markdown
from .projeto import Projeto, slug

FORMATOS = ("epub", "docx", "md", "html")


def _esc(t: str) -> str:
    return html.escape(t, quote=False)


# ----------------------------------------------------------------------------- HTML lado a lado

_CSS_LADO_A_LADO = """
:root{color-scheme:light dark;--bg:#fff;--fg:#1c1c1e;--mut:#6b6b70;--bd:#d9d9de;--ac:#7a3e9d;--bad:#b3261e}
@media (prefers-color-scheme:dark){:root{--bg:#161618;--fg:#e8e8ea;--mut:#9a9aa0;--bd:#38383d;--ac:#c79bff;--bad:#ff8a80}}
body{margin:0;background:var(--bg);color:var(--fg);font:16px/1.55 Georgia,serif}
main{max-width:1280px;margin:0 auto;padding:24px 20px 80px}
h1{font:600 22px/1.3 system-ui,sans-serif;margin:0 0 4px}
.meta{color:var(--mut);font:13px/1.5 system-ui,sans-serif;margin-bottom:20px}
table{width:100%;border-collapse:collapse;table-layout:fixed}
th{font:600 12px system-ui,sans-serif;text-transform:uppercase;letter-spacing:.06em;color:var(--mut);text-align:left;
   padding:8px 10px;border-bottom:2px solid var(--bd);position:sticky;top:0;background:var(--bg)}
td{vertical-align:top;padding:10px;border-bottom:1px solid var(--bd);overflow-wrap:anywhere}
td.n{width:42px;color:var(--mut);font:12px system-ui,sans-serif}
td.titulo{font-weight:700;font-size:1.1em} td.verso{font-style:italic} td.nota{font-size:.88em;color:var(--mut)}
.badge{display:inline-block;margin-left:6px;padding:0 6px;border-radius:9px;border:1px solid var(--bad);
   color:var(--bad);font:11px system-ui,sans-serif}
.badge.info{border-color:var(--mut);color:var(--mut)}
@media (max-width:800px){td.n{display:none}}
"""


def pagina_html(titulo: str, subtitulo: str, linhas: list[dict], rotulos: tuple[str, ...]) -> str:
    """linhas: [{n, tipo, orig, trad, base (opcional), alertas}]; 3 colunas se houver 'base'."""
    tem_base = any(l.get("base") is not None for l in linhas)
    cab = "<th></th>" + "".join(f"<th>{_esc(r)}</th>" for r in rotulos[: 3 if tem_base else 2])

    def cel(texto: str, tipo: str) -> str:
        return f'<td class="{tipo}">' + "<br>".join(_esc(x) for x in texto.split("\n")) + "</td>"

    corpo = []
    for l in linhas:
        selos = "".join(
            f'<span class="badge{" info" if a in ("corrigido_manualmente", "ja_no_idioma_de_destino") else ""}">'
            f"{_esc(a)}</span>"
            for a in l.get("alertas", [])
        )
        trad = cel(l["trad"], l["tipo"])[:-5] + selos + "</td>"
        base = cel(l["base"], l["tipo"]) if tem_base and l.get("base") is not None else ("<td></td>" if tem_base else "")
        corpo.append(f'<tr><td class="n">{l["n"]}</td>{cel(l["orig"], l["tipo"])}{trad}{base}</tr>')
    return (
        '<!doctype html><html lang="pt"><head><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width,initial-scale=1">'
        f"<title>{_esc(titulo)}</title><style>{_CSS_LADO_A_LADO}</style></head><body><main>"
        f'<h1>{_esc(titulo)}</h1><div class="meta">{subtitulo}</div>'
        f"<table><thead><tr>{cab}</tr></thead><tbody>{''.join(corpo)}</tbody></table></main></body></html>"
    )


# ----------------------------------------------------------------------------- EPUB

_CSS_EPUB = """
body{font-family:serif;line-height:1.5;margin:5%}
h1,h2,h3,h4,h5,h6{font-family:sans-serif;line-height:1.2}
h1{margin:2em 0 1em}
p{margin:0;text-indent:1.4em;text-align:justify}
h1+p,h2+p,h3+p,h4+p,hr+p,p.verso+p{text-indent:0}
p.verso{text-indent:0;margin:1em 0 1em 2em;font-style:italic;text-align:left}
p.nota{font-size:.85em;text-indent:0;margin:.5em 0;border-top:1px solid #999;padding-top:.3em}
p.legenda{text-align:center;font-size:.9em;font-style:italic;text-indent:0}
hr{border:0;text-align:center;margin:1.5em 0}
hr:after{content:"* * *"}
"""


def _lang(codigo: str) -> str:
    return {"zh-cn": "zh-CN", "zh-tw": "zh-TW"}.get(codigo, codigo or "und")


def _html_dos_blocos(blocos: list[Bloco], textos: dict[str, str]) -> str:
    niveis = [b.nivel for b in blocos if b.tipo == TITULO]
    base = min(niveis) if niveis else 1
    saida = []
    for b in blocos:
        t = textos.get(b.id, b.texto)
        if b.tipo == TITULO:
            n = min(6, max(1, b.nivel - base + 1))
            saida.append(f"<h{n}>{_esc(t)}</h{n}>")
        elif b.tipo == VERSO:
            saida.append('<p class="verso">' + "<br/>\n".join(_esc(l) for l in t.split("\n")) + "</p>")
        elif b.tipo == NOTA:
            saida.append(f'<p class="nota" epub:type="footnote">{_esc(t)}</p>')
        elif b.tipo == LEGENDA:
            saida.append(f'<p class="legenda">{_esc(t)}</p>')
        elif b.tipo == SEPARADOR:
            saida.append("<hr/>")
        else:
            saida.append(f"<p>{_esc(t)}</p>")
    return "\n".join(saida)


def escrever_epub(destino: Path, titulo: str, autor: str, idioma: str, capitulos: list[tuple[str, str]]) -> None:
    uid = "urn:uuid:" + str(uuid.uuid4())
    agora = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    lang = _lang(idioma)
    manifesto = [
        '<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>',
        '<item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/>',
        '<item id="css" href="style.css" media-type="text/css"/>',
    ]
    spine, nav, ncx, arquivos = [], [], [], []
    for i, (tit, corpo) in enumerate(capitulos, 1):
        nome = f"cap{i:03d}.xhtml"
        manifesto.append(f'<item id="c{i}" href="{nome}" media-type="application/xhtml+xml"/>')
        spine.append(f'<itemref idref="c{i}"/>')
        nav.append(f'<li><a href="{nome}">{_esc(tit)}</a></li>')
        ncx.append(
            f'<navPoint id="n{i}" playOrder="{i}"><navLabel><text>{_esc(tit)}</text></navLabel>'
            f'<content src="{nome}"/></navPoint>'
        )
        arquivos.append((nome, (
            '<?xml version="1.0" encoding="utf-8"?>\n<!DOCTYPE html>\n'
            f'<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" '
            f'xml:lang="{lang}" lang="{lang}"><head><meta charset="utf-8"/><title>{_esc(tit)}</title>'
            f'<link rel="stylesheet" type="text/css" href="style.css"/></head><body>\n{corpo}\n</body></html>'
        )))
    criador = f"<dc:creator>{_esc(autor)}</dc:creator>" if autor else ""
    opf = (
        '<?xml version="1.0" encoding="utf-8"?>\n'
        f'<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="bookid" xml:lang="{lang}">'
        '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">'
        f'<dc:identifier id="bookid">{uid}</dc:identifier><dc:title>{_esc(titulo)}</dc:title>{criador}'
        f'<dc:language>{lang}</dc:language><meta property="dcterms:modified">{agora}</meta></metadata>'
        f'<manifest>{"".join(manifesto)}</manifest><spine toc="ncx">{"".join(spine)}</spine></package>'
    )
    nav_xhtml = (
        '<?xml version="1.0" encoding="utf-8"?>\n<!DOCTYPE html>\n'
        f'<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" lang="{lang}">'
        f'<head><meta charset="utf-8"/><title>{_esc(titulo)}</title></head><body>'
        f'<nav epub:type="toc" id="toc"><h1>{_esc(titulo)}</h1><ol>{"".join(nav)}</ol></nav></body></html>'
    )
    ncx_xml = (
        '<?xml version="1.0" encoding="UTF-8"?>\n<ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" version="2005-1">'
        f'<head><meta name="dtb:uid" content="{uid}"/><meta name="dtb:depth" content="1"/>'
        '<meta name="dtb:totalPageCount" content="0"/><meta name="dtb:maxPageNumber" content="0"/></head>'
        f'<docTitle><text>{_esc(titulo)}</text></docTitle><navMap>{"".join(ncx)}</navMap></ncx>'
    )
    container = (
        '<?xml version="1.0"?>\n<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
        '<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>'
        "</rootfiles></container>"
    )
    with zipfile.ZipFile(destino, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        z.writestr("META-INF/container.xml", container)
        z.writestr("OEBPS/content.opf", opf)
        z.writestr("OEBPS/nav.xhtml", nav_xhtml)
        z.writestr("OEBPS/toc.ncx", ncx_xml)
        z.writestr("OEBPS/style.css", _CSS_EPUB)
        for nome, conteudo in arquivos:
            z.writestr("OEBPS/" + nome, conteudo)


# ----------------------------------------------------------------------------- DOCX

def escrever_docx(destino: Path, titulo: str, autor: str, livro: Livro, textos: dict[str, str]) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Document()
    doc.core_properties.title = titulo
    doc.core_properties.author = autor or ""
    doc.add_heading(titulo, 0)
    if autor:
        doc.add_paragraph(autor)
    for ci, cap in enumerate(livro.capitulos):
        if ci > 0 or titulo:
            doc.add_page_break()
        niveis = [b.nivel for b in cap.blocos if b.tipo == TITULO]
        base = min(niveis) if niveis else 1
        for b in cap.blocos:
            t = textos.get(b.id, b.texto)
            if b.tipo == TITULO:
                doc.add_heading(t, level=min(4, max(1, b.nivel - base + 1)))
            elif b.tipo == VERSO:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Pt(28)
                for i, linha in enumerate(t.split("\n")):
                    run = p.add_run(linha)
                    run.italic = True
                    if i < len(t.split("\n")) - 1:
                        run.add_break()
            elif b.tipo == NOTA:
                run = doc.add_paragraph().add_run(t)
                run.font.size = Pt(9)
                run.italic = True
            elif b.tipo == LEGENDA:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run(t).italic = True
            elif b.tipo == SEPARADOR:
                p = doc.add_paragraph("* * *")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                doc.add_paragraph(t)
    doc.save(str(destino))


# ----------------------------------------------------------------------------- orquestração

def _titulo_traduzido(livro: Livro, textos: dict[str, str]) -> str:
    alvo = (livro.titulo or "").strip().lower()
    for b in livro.blocos():
        if b.tipo == TITULO and b.texto.strip().lower() == alvo:
            return textos.get(b.id, livro.titulo)
    return livro.titulo


def exportar_tudo(proj: Projeto, cfg: Config, livro: Livro, textos: dict[str, str], resultados: dict,
                  formatos: list[str]) -> list[Path]:
    saida = proj.pasta("saida")
    base = f"{slug(livro.titulo or 'livro')}.{cfg.destino}"
    titulo = _titulo_traduzido(livro, textos) or livro.titulo
    gerados: list[Path] = []

    if "md" in formatos:
        d = Livro(titulo=titulo, autor=livro.autor, idioma=cfg.destino, capitulos=livro.capitulos)
        caminho = saida / f"{base}.md"
        caminho.write_text(livro_para_markdown(d, textos), encoding="utf-8")
        gerados.append(caminho)

    if "epub" in formatos:
        capitulos = []
        for i, cap in enumerate(livro.capitulos, 1):
            tit = next((textos.get(b.id, b.texto) for b in cap.blocos if b.tipo == TITULO), f"Parte {i}")
            capitulos.append((tit, _html_dos_blocos(cap.blocos, textos)))
        caminho = saida / f"{base}.epub"
        escrever_epub(caminho, titulo, livro.autor, cfg.destino, capitulos)
        gerados.append(caminho)

    if "docx" in formatos:
        caminho = saida / f"{base}.docx"
        escrever_docx(caminho, titulo, livro.autor, livro, textos)
        gerados.append(caminho)

    if "html" in formatos:
        linhas = []
        for n, b in enumerate(livro.blocos(), 1):
            r = resultados.get(b.id)
            linhas.append({"n": n, "tipo": b.tipo, "orig": b.texto, "trad": textos.get(b.id, b.texto),
                           "alertas": r.alertas if r else []})
        subtitulo = (f"{_esc(nome_idioma(cfg.origem))} → {_esc(nome_idioma(cfg.destino))} · modelo "
                     f"{_esc(cfg.modelo)} · {len(linhas)} trechos")
        caminho = saida / f"{base}.bilingue.html"
        caminho.write_text(
            pagina_html(f"{livro.titulo} — edição bilíngue", subtitulo, linhas,
                        (nome_idioma(cfg.origem).capitalize(), nome_idioma(cfg.destino).capitalize())),
            encoding="utf-8",
        )
        gerados.append(caminho)
    return gerados
