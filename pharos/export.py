"""Export: Markdown, EPUB, DOCX and bilingual HTML (EPUB written with zipfile, no dependencies)."""
from __future__ import annotations

import html
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from .config import Config, language_name
from .model import Block, Book, CAPTION, NOTE, SEPARATOR, TITLE, VERSE, book_to_markdown
from .project import Project, slug

FORMATS = ("epub", "docx", "md", "html")


def _esc(t: str) -> str:
    return html.escape(t, quote=False)


# ----------------------------------------------------------------------------- side-by-side HTML

_CSS_SIDE_BY_SIDE = """
:root{color-scheme:light dark;--bg:#fff;--fg:#1c1c1e;--mut:#6b6b70;--bd:#d9d9de;--ac:#7a3e9d;--bad:#b3261e}
@media (prefers-color-scheme:dark){:root{--bg:#161618;--fg:#e8e8ea;--mut:#9a9aa0;--bd:#38383d;--ac:#c79bff;--bad:#ff8a80}}
body{margin:0;background:var(--bg);color:var(--fg);font:16px/1.55 Georgia,serif}
main{max-width:1280px;margin:0 auto;padding:24px 20px 80px}
h1{font:600 22px/1.3 system-ui,sans-serif;margin:0 0 4px}
.meta{color:var(--mut);font:13px/1.5 system-ui,sans-serif;margin-bottom:20px}
table{width:100%;border-collapse:collapse;table-layout:fixed}
th{font:600 12px system-ui,sans-serif;text-transform:uppercase;letter-spacing:.06em;color:var(--mut);text-align:left;
   padding:8px 10px;border-bottom:2px solid var(--bd);position:sticky;top:0;background:var(--bg)}
td{vertical-align:top;padding:10px;border-bottom:1px solid var(--bd);overflow-wrap:anywhere}
td.n{width:42px;color:var(--mut);font:12px system-ui,sans-serif}
td.title{font-weight:700;font-size:1.1em} td.verse{font-style:italic} td.note{font-size:.88em;color:var(--mut)}
.badge{display:inline-block;margin-left:6px;padding:0 6px;border-radius:9px;border:1px solid var(--bad);
   color:var(--bad);font:11px system-ui,sans-serif}
.badge.info{border-color:var(--mut);color:var(--mut)}
@media (max-width:800px){td.n{display:none}}
"""


def page_html(title: str, subtitle: str, lines: list[dict], labels: tuple[str, ...]) -> str:
    """lines: [{n, kind, orig, trans, base (optional), alerts}]; 3 columns if 'base' is present."""
    has_base = any(l.get("base") is not None for l in lines)
    head = "<th></th>" + "".join(f"<th>{_esc(r)}</th>" for r in labels[: 3 if has_base else 2])

    def cell(text: str, kind: str) -> str:
        return f'<td class="{kind}">' + "<br>".join(_esc(x) for x in text.split("\n")) + "</td>"

    body = []
    for l in lines:
        badges = "".join(
            f'<span class="badge{" info" if a in ("manually_corrected", "already_in_target") else ""}">'
            f"{_esc(a)}</span>"
            for a in l.get("alerts", [])
        )
        trans = cell(l["trans"], l["kind"])[:-5] + badges + "</td>"
        base = cell(l["base"], l["kind"]) if has_base and l.get("base") is not None else ("<td></td>" if has_base else "")
        body.append(f'<tr><td class="n">{l["n"]}</td>{cell(l["orig"], l["kind"])}{trans}{base}</tr>')
    return (
        '<!doctype html><html lang="pt"><head><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width,initial-scale=1">'
        f"<title>{_esc(title)}</title><style>{_CSS_SIDE_BY_SIDE}</style></head><body><main>"
        f'<h1>{_esc(title)}</h1><div class="meta">{subtitle}</div>'
        f"<table><thead><tr>{head}</tr></thead><tbody>{''.join(body)}</tbody></table></main></body></html>"
    )


# ----------------------------------------------------------------------------- EPUB

_CSS_EPUB = """
body{font-family:serif;line-height:1.5;margin:5%}
h1,h2,h3,h4,h5,h6{font-family:sans-serif;line-height:1.2}
h1{margin:2em 0 1em}
p{margin:0;text-indent:1.4em;text-align:justify}
h1+p,h2+p,h3+p,h4+p,hr+p,p.verse+p{text-indent:0}
p.verse{text-indent:0;margin:1em 0 1em 2em;font-style:italic;text-align:left}
p.note{font-size:.85em;text-indent:0;margin:.5em 0;border-top:1px solid #999;padding-top:.3em}
p.caption{text-align:center;font-size:.9em;font-style:italic;text-indent:0}
hr{border:0;text-align:center;margin:1.5em 0}
hr:after{content:"* * *"}
"""


def _lang(code: str) -> str:
    return {"zh-cn": "zh-CN", "zh-tw": "zh-TW"}.get(code, code or "und")


def _blocks_html(blocks: list[Block], texts: dict[str, str]) -> str:
    levels = [b.level for b in blocks if b.kind == TITLE]
    base = min(levels) if levels else 1
    out = []
    for b in blocks:
        t = texts.get(b.id, b.text)
        if b.kind == TITLE:
            n = min(6, max(1, b.level - base + 1))
            out.append(f"<h{n}>{_esc(t)}</h{n}>")
        elif b.kind == VERSE:
            out.append('<p class="verse">' + "<br/>\n".join(_esc(l) for l in t.split("\n")) + "</p>")
        elif b.kind == NOTE:
            out.append(f'<p class="note" epub:type="footnote">{_esc(t)}</p>')
        elif b.kind == CAPTION:
            out.append(f'<p class="caption">{_esc(t)}</p>')
        elif b.kind == SEPARATOR:
            out.append("<hr/>")
        else:
            out.append(f"<p>{_esc(t)}</p>")
    return "\n".join(out)


def write_epub(dest: Path, title: str, author: str, language: str, chapters: list[tuple[str, str]]) -> None:
    uid = "urn:uuid:" + str(uuid.uuid4())
    now = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    lang = _lang(language)
    manifest = [
        '<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>',
        '<item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/>',
        '<item id="css" href="style.css" media-type="text/css"/>',
    ]
    spine, nav, ncx, files = [], [], [], []
    for i, (ttl, body) in enumerate(chapters, 1):
        name = f"cap{i:03d}.xhtml"
        manifest.append(f'<item id="c{i}" href="{name}" media-type="application/xhtml+xml"/>')
        spine.append(f'<itemref idref="c{i}"/>')
        nav.append(f'<li><a href="{name}">{_esc(ttl)}</a></li>')
        ncx.append(
            f'<navPoint id="n{i}" playOrder="{i}"><navLabel><text>{_esc(ttl)}</text></navLabel>'
            f'<content src="{name}"/></navPoint>'
        )
        files.append((name, (
            '<?xml version="1.0" encoding="utf-8"?>\n<!DOCTYPE html>\n'
            f'<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" '
            f'xml:lang="{lang}" lang="{lang}"><head><meta charset="utf-8"/><title>{_esc(ttl)}</title>'
            f'<link rel="stylesheet" type="text/css" href="style.css"/></head><body>\n{body}\n</body></html>'
        )))
    creator = f"<dc:creator>{_esc(author)}</dc:creator>" if author else ""
    opf = (
        '<?xml version="1.0" encoding="utf-8"?>\n'
        f'<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="bookid" xml:lang="{lang}">'
        '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">'
        f'<dc:identifier id="bookid">{uid}</dc:identifier><dc:title>{_esc(title)}</dc:title>{creator}'
        f'<dc:language>{lang}</dc:language><meta property="dcterms:modified">{now}</meta></metadata>'
        f'<manifest>{"".join(manifest)}</manifest><spine toc="ncx">{"".join(spine)}</spine></package>'
    )
    nav_xhtml = (
        '<?xml version="1.0" encoding="utf-8"?>\n<!DOCTYPE html>\n'
        f'<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" lang="{lang}">'
        f'<head><meta charset="utf-8"/><title>{_esc(title)}</title></head><body>'
        f'<nav epub:type="toc" id="toc"><h1>{_esc(title)}</h1><ol>{"".join(nav)}</ol></nav></body></html>'
    )
    ncx_xml = (
        '<?xml version="1.0" encoding="UTF-8"?>\n<ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" version="2005-1">'
        f'<head><meta name="dtb:uid" content="{uid}"/><meta name="dtb:depth" content="1"/>'
        '<meta name="dtb:totalPageCount" content="0"/><meta name="dtb:maxPageNumber" content="0"/></head>'
        f'<docTitle><text>{_esc(title)}</text></docTitle><navMap>{"".join(ncx)}</navMap></ncx>'
    )
    container = (
        '<?xml version="1.0"?>\n<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
        '<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>'
        "</rootfiles></container>"
    )
    with zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        z.writestr("META-INF/container.xml", container)
        z.writestr("OEBPS/content.opf", opf)
        z.writestr("OEBPS/nav.xhtml", nav_xhtml)
        z.writestr("OEBPS/toc.ncx", ncx_xml)
        z.writestr("OEBPS/style.css", _CSS_EPUB)
        for name, content in files:
            z.writestr("OEBPS/" + name, content)


# ----------------------------------------------------------------------------- DOCX

def write_docx(dest: Path, title: str, author: str, book: Book, texts: dict[str, str]) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Document()
    doc.core_properties.title = title
    doc.core_properties.author = author or ""
    doc.add_heading(title, 0)
    if author:
        doc.add_paragraph(author)
    for ci, ch in enumerate(book.chapters):
        if ci > 0 or title:
            doc.add_page_break()
        levels = [b.level for b in ch.blocks if b.kind == TITLE]
        base = min(levels) if levels else 1
        for b in ch.blocks:
            t = texts.get(b.id, b.text)
            if b.kind == TITLE:
                doc.add_heading(t, level=min(4, max(1, b.level - base + 1)))
            elif b.kind == VERSE:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Pt(28)
                for i, line in enumerate(t.split("\n")):
                    run = p.add_run(line)
                    run.italic = True
                    if i < len(t.split("\n")) - 1:
                        run.add_break()
            elif b.kind == NOTE:
                run = doc.add_paragraph().add_run(t)
                run.font.size = Pt(9)
                run.italic = True
            elif b.kind == CAPTION:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run(t).italic = True
            elif b.kind == SEPARATOR:
                p = doc.add_paragraph("* * *")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                doc.add_paragraph(t)
    doc.save(str(dest))


# ----------------------------------------------------------------------------- orchestration

def _translated_title(book: Book, texts: dict[str, str]) -> str:
    target = (book.title or "").strip().lower()
    for b in book.blocks():
        if b.kind == TITLE and b.text.strip().lower() == target:
            return texts.get(b.id, book.title)
    return book.title


def export_all(proj: Project, cfg: Config, book: Book, texts: dict[str, str], results: dict,
               formats: list[str]) -> list[Path]:
    out_dir = proj.folder("output")
    base = f"{slug(book.title or 'book')}.{cfg.target}"
    title = _translated_title(book, texts) or book.title
    generated: list[Path] = []

    if "md" in formats:
        d = Book(title=title, author=book.author, language=cfg.target, chapters=book.chapters)
        path = out_dir / f"{base}.md"
        path.write_text(book_to_markdown(d, texts), encoding="utf-8")
        generated.append(path)

    if "epub" in formats:
        chapters = []
        for i, ch in enumerate(book.chapters, 1):
            ttl = next((texts.get(b.id, b.text) for b in ch.blocks if b.kind == TITLE), f"Part {i}")
            chapters.append((ttl, _blocks_html(ch.blocks, texts)))
        path = out_dir / f"{base}.epub"
        write_epub(path, title, book.author, cfg.target, chapters)
        generated.append(path)

    if "docx" in formats:
        path = out_dir / f"{base}.docx"
        write_docx(path, title, book.author, book, texts)
        generated.append(path)

    if "html" in formats:
        lines = []
        for n, b in enumerate(book.blocks(), 1):
            r = results.get(b.id)
            lines.append({"n": n, "kind": b.kind, "orig": b.text, "trans": texts.get(b.id, b.text),
                          "alerts": r.alerts if r else []})
        subtitle = (f"{_esc(language_name(cfg.source))} → {_esc(language_name(cfg.target))} · model "
                    f"{_esc(cfg.model)} · {len(lines)} passages")
        path = out_dir / f"{base}.bilingual.html"
        path.write_text(
            page_html(f"{book.title} — bilingual edition", subtitle, lines,
                      (language_name(cfg.source).capitalize(), language_name(cfg.target).capitalize())),
            encoding="utf-8",
        )
        generated.append(path)
    return generated
